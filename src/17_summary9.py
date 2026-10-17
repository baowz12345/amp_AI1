# -*- coding: utf-8 -*-
"""
17_summary9.py   九套方案汇总: 复用已实现的 7 方案结果(S1/S2/S7) + 本轮新增
================
- 读取 results/metrics_*.csv 的 pooled_OOF 行
- 映射为 9 套方案(用户表格口径), 输出:
    results/comparison_schemes9.csv / .xlsx(两 sheet: 9方案汇总 + 18模型明细)
    results/figures/scheme9_performance.png (AUC/F1/MCC 柱状)
    results/figures/scheme9_roc_pr.png (9 代表模型 OOF ROC/PR 曲线)
"""
import os
import sys
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from sklearn.metrics import roc_curve, precision_recall_curve, auc

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common as C

# 方案定义: 编号, 名称, 输入, 亮点, 算力, 可解释性, 候选模型名(按优先级取最优)
SCHEMES = [
    (1, "手工特征 + 传统ML", "理化、AAC、k-mer", "基线，可解释", "低", "⭐⭐⭐⭐⭐",
     ["S2_XGB", "S1_XGB", "S2_RF", "S1_RF"]),
    (2, "CNN-BiLSTM 端到端", "氨基酸序列", "自动学习motif", "中", "⭐⭐",
     ["S8_CNNBiLSTM"]),
    (3, "ProtBERT 微调", "氨基酸序列", "SOTA基线，迁移学习", "高", "⭐",
     ["S9_ProtBERT_SVM", "S9_ProtBERT_LR"]),
    (4, "GAT/GCN 图学习", "氨基酸图+节点理化特征", "建模残基相互作用", "中高", "⭐⭐",
     ["S10_GCN", "S10_GAT"]),
    (5, "对比自监督预训练", "无标注肽序列+标注APD", "小样本友好", "极高", "⭐",
     ["S11_Contrastive_SVM", "S11_Contrastive_LR"]),
    (6, "TF-IDF k-mer + SVM/LR", "k-mer词袋", "最简基线", "极低", "⭐⭐⭐",
     ["S12_TFIDF_SVM", "S12_TFIDF_LR"]),
    (7, "多模型集成 Stacking", "多模型概率输出", "整体性能上限", "很高", "⭐",
     ["S7_stackLR", "S7_stackXGB"]),
    (8, "MoE 混合专家", "序列+理化", "门控按肽长度分治", "中高", "⭐⭐",
     ["S13_MoE"]),
    (9, "序列+二级结构多模态", "序列+预测二级结构", "引入结构生物学先验", "高", "⭐⭐",
     ["S14_MultiModal"]),
]
METRIC_COLS = ["TP", "FP", "TN", "FN", "Acc", "Precision", "Recall", "F1", "MCC", "AUC", "PR_AUC"]


def load_pooled():
    rows = []
    for f in os.listdir(C.RESULT_DIR):
        if f.startswith("metrics_") and f.endswith(".csv"):
            name = f[len("metrics_"):-len(".csv")]
            df = pd.read_csv(os.path.join(C.RESULT_DIR, f))
            pooled = df[df["fold"] == "pooled_OOF"]
            if pooled.empty:
                continue
            row = pooled.iloc[0].to_dict()
            row["model"] = name
            rows.append(row)
    return pd.DataFrame(rows)


def main():
    C.setup_chinese_font()
    pooled_df = load_pooled()
    # 9 方案汇总(每方案取候选模型中 AUC 最优且在池中存在者)
    summary = []
    for (num, sname, inp, hl, cost, expl, cands) in SCHEMES:
        avail = [m for m in cands if m in set(pooled_df["model"])]
        if not avail:
            continue
        best = max(avail, key=lambda m: pooled_df.loc[pooled_df["model"] == m, "AUC"].iloc[0])
        row = pooled_df[pooled_df["model"] == best].iloc[0].to_dict()
        summary.append({"编号": num, "方案": sname, "输入": inp, "核心亮点": hl,
                        "算力": cost, "可解释性": expl, "代表模型": best,
                        **{c: row[c] for c in METRIC_COLS}})
    df9 = pd.DataFrame(summary)
    df9.to_csv(os.path.join(C.RESULT_DIR, "comparison_schemes9.csv"),
               index=False, encoding="utf-8-sig")

    # xlsx 两 sheet
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "9方案汇总"
    header_fill = PatternFill("solid", fgColor="D9D9D9")
    for j, col in enumerate(df9.columns, 1):
        c = ws.cell(row=1, column=j, value=col)
        c.font = Font(bold=True)
        c.fill = header_fill
    for i, row in df9.iterrows():
        for j, col in enumerate(df9.columns, 1):
            v = row[col]
            if isinstance(v, (float, np.floating)):
                v = round(float(v), 4)
            ws.cell(row=i + 2, column=j, value=v)
    for j in range(1, len(df9.columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 16
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("18模型明细")
    mlist = pooled_df[["model"] + METRIC_COLS].sort_values("AUC", ascending=False).reset_index(drop=True)
    for j, col in enumerate(mlist.columns, 1):
        c = ws2.cell(row=1, column=j, value=col)
        c.font = Font(bold=True)
        c.fill = header_fill
    for i, row in mlist.iterrows():
        for j, col in enumerate(mlist.columns, 1):
            v = row[col]
            if isinstance(v, (float, np.floating)):
                v = round(float(v), 4)
            ws2.cell(row=i + 2, column=j, value=v)
    ws2.freeze_panes = "A2"
    for j in range(1, len(mlist.columns) + 1):
        ws2.column_dimensions[get_column_letter(j)].width = 16
    wb.save(os.path.join(C.RESULT_DIR, "comparison_schemes9.xlsx"))

    # 柱状图
    d = df9.sort_values("AUC", ascending=False)
    x = np.arange(len(d))
    w = 0.26
    fig, ax = plt.subplots(figsize=(12, 6), dpi=150)
    ax.bar(x - w, d["AUC"], w, label="AUC", color="#9BBBF4")
    ax.bar(x, d["F1"], w, label="F1", color="#A3D5E8")
    ax.bar(x + w, d["MCC"], w, label="MCC", color="#94D8C3")
    ax.set_xticks(x)
    ax.set_xticklabels([f"{n}.{s}" for n, s in zip(d["编号"], d["方案"])], rotation=30, ha="right", fontsize=9)
    ax.set_ylim(0.4, 1.0)
    ax.set_ylabel("Score")
    ax.set_title("九套方案性能对比 (10折CV pooled-OOF)")
    ax.legend()
    ax.grid(axis="y", alpha=0.3)
    fig.tight_layout()
    fig.savefig(os.path.join(C.FIG_DIR, "scheme9_performance.png"))
    plt.close(fig)

    # ROC/PR 图(9 代表模型)
    y = C.load_dataset()["label"].values
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(13, 5.5), dpi=150)
    colors = plt.cm.tab10(np.linspace(0, 1, 10))
    for k, row in df9.iterrows():
        m = row["代表模型"]
        path = os.path.join(C.RESULT_DIR, f"oof_{m}.npy")
        if not os.path.exists(path):
            continue
        prob = np.load(path)
        fpr, tpr, _ = roc_curve(y, prob)
        pr, rc, _ = precision_recall_curve(y, prob)
        lab = f"{int(row['编号'])}.{row['方案']}"
        ax1.plot(fpr, tpr, color=colors[k % 10], lw=1.6, label=f"{lab} (AUC={row['AUC']:.3f})")
        ax2.plot(rc, pr, color=colors[k % 10], lw=1.6,
                 label=f"{lab} (PR={row['PR_AUC']:.3f})")
    ax1.plot([0, 1], [0, 1], "k--", lw=0.8, alpha=0.5)
    ax1.set_xlabel("FPR"); ax1.set_ylabel("TPR"); ax1.set_title("ROC (pooled OOF)")
    ax1.legend(fontsize=7, loc="lower right")
    ax2.set_xlabel("Recall"); ax2.set_ylabel("Precision"); ax2.set_title("PR (pooled OOF)")
    ax2.legend(fontsize=7, loc="lower left")
    fig.tight_layout()
    fig.savefig(os.path.join(C.FIG_DIR, "scheme9_roc_pr.png"))
    plt.close(fig)

    print("=" * 70)
    print(df9[["编号", "方案", "代表模型", "AUC", "PR_AUC", "MCC", "Acc", "F1"]].to_string(index=False))
    print("=" * 70)
    print("已输出: comparison_schemes9.csv/.xlsx, figures/scheme9_performance.png, scheme9_roc_pr.png")


if __name__ == "__main__":
    main()
